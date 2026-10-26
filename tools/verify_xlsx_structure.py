#!/usr/bin/env python3
"""Verify the retained XLSX sample with an independent openpyxl reader."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import warnings

from openpyxl import load_workbook
import openpyxl


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def check_cell(worksheet, expected: dict[str, object]) -> None:
    coordinate = str(expected["coordinate"])
    cell = worksheet[coordinate]
    actual = {"type": cell.data_type, "value": cell.value}
    wanted = {"type": expected["type"], "value": expected["value"]}
    if actual != wanted:
        raise AssertionError(
            f"{worksheet.title}!{coordinate}: expected {wanted!r}, got {actual!r}"
        )


def verify(sample: Path, expectations_path: Path) -> dict[str, object]:
    expectations = json.loads(expectations_path.read_text(encoding="utf-8"))
    with warnings.catch_warnings(record=True) as caught_warnings:
        warnings.simplefilter("always")
        workbook = load_workbook(sample, data_only=False, read_only=False)
    style_warnings = [
        str(warning.message)
        for warning in caught_warnings
        if "style" in str(warning.message).lower()
    ]
    if style_warnings:
        workbook.close()
        raise AssertionError(f"openpyxl reported style warnings: {style_warnings!r}")
    try:
        expected_sheetnames = expectations["sheetnames"]
        if workbook.sheetnames != expected_sheetnames:
            raise AssertionError(
                f"sheet order: expected {expected_sheetnames!r}, "
                f"got {workbook.sheetnames!r}"
            )

        wide_spec = expectations["wide_sheet"]
        wide = workbook[expected_sheetnames[0]]
        if wide.max_column != wide_spec["header_count"]:
            raise AssertionError(
                f"wide sheet max column: expected {wide_spec['header_count']}, "
                f"got {wide.max_column}"
            )
        header_row = wide_spec["header_row"]
        expected_headers = [
            f"{wide_spec['header_prefix']}{index}"
            for index in range(1, wide_spec["header_count"] + 1)
        ]
        actual_headers = [
            wide.cell(row=header_row, column=index).value
            for index in range(1, wide_spec["header_count"] + 1)
        ]
        if actual_headers != expected_headers:
            raise AssertionError("wide sheet column headers do not match the oracle")

        for expected in wide_spec["cells"]:
            check_cell(wide, expected)

        second = workbook[expected_sheetnames[1]]
        for expected in expectations["second_sheet_cells"]:
            check_cell(second, expected)

        formula_cells = [
            f"{worksheet.title}!{cell.coordinate}"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        if formula_cells:
            raise AssertionError(f"formula-like text became formulas: {formula_cells!r}")

        return {
            "status": "PASS",
            "sample": str(sample),
            "sample_sha256": sha256(sample),
            "expectations": str(expectations_path),
            "expectations_sha256": sha256(expectations_path),
            "openpyxl_version": openpyxl.__version__,
            "sheets": workbook.sheetnames,
            "wide_column_count": wide.max_column,
            "checked_cells": len(wide_spec["cells"])
            + len(expectations["second_sheet_cells"]),
            "style_warnings": style_warnings,
            "formula_cells": formula_cells,
            "nonfinite_inputs_rejected_by_unit_test": expectations[
                "nonfinite_inputs_rejected_by_unit_test"
            ],
            "determinism_checked_by_unit_test": expectations[
                "determinism_checked_by_unit_test"
            ],
        }
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("sample", type=Path, help="generated .xlsx workbook")
    parser.add_argument("expectations", type=Path, help="oracle JSON file")
    args = parser.parse_args()
    print(json.dumps(verify(args.sample, args.expectations), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
